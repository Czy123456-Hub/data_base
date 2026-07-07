from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook


SOURCE_DOCUMENT = "中国港口信息.xlsx"
MODULE_SLUG = "port_agency_info"


def one_line(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def content_lines(value: object) -> list[str]:
    return [one_line(line) for line in str(value or "").splitlines() if one_line(line)]


def clean_raw_text(value: object) -> str:
    return "\n".join(line.rstrip() for line in str(value or "").strip().splitlines())


def is_address_line(line: str) -> bool:
    return re.search(r"^(add|address)\b\s*[:：]?", line, re.I) is not None


def is_fax_line(line: str) -> bool:
    return re.search(r"^(fax|facsimile|tel\s*/\s*fax|tel\s*&\s*fax)\b\s*[:：]?", line, re.I) is not None


def is_contact_line(line: str) -> bool:
    return re.search(
        r"\b(pic|attn|contact|operator|manager|mr\.?|ms\.?|mrs\.?|wechat|we\s*chat|op)\b",
        line,
        re.I,
    ) is not None


def is_tel_line(line: str) -> bool:
    if is_contact_line(line) or is_fax_line(line):
        return False
    return re.search(
        r"^(tel|telephone|tele\s*no|office\s*tel|office\s*phone|phone|mobile|mob|cell\s*phone|direct\s*line|mb)\b\s*[:：]?",
        line,
        re.I,
    ) is not None


def clean_label(line: str) -> str:
    text = one_line(line)
    patterns = [
        r"^(add|address)\b\s*[:：]?",
        r"^(fax|facsimile|tel\s*/\s*fax|tel\s*&\s*fax)\b\s*[:：]?",
        r"^(tel|telephone|tele\s*no|office\s*tel|office\s*phone|phone|mobile|mob|cell\s*phone|direct\s*line|mb)\b\s*[:：]?",
        r"^(e-?mail|email)\b\s*[:：]?",
    ]
    for pattern in patterns:
        text = re.sub(pattern, "", text, flags=re.I).strip()
    return text


def unique(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        item = one_line(value)
        if item and item not in seen:
            seen.add(item)
            result.append(item)
    return result


def parse_agent(raw_text: str) -> dict[str, str | None]:
    lines = content_lines(raw_text)
    emails: list[str] = []
    for line in lines:
        emails.extend(re.findall(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", line, flags=re.I))

    parsed = {
        "agency_name": lines[0] if lines else None,
        "address": " / ".join(unique([clean_label(line) for line in lines if is_address_line(line)])),
        "tel": " / ".join(unique([clean_label(line) for line in lines if is_tel_line(line)])),
        "fax": " / ".join(unique([clean_label(line) for line in lines if is_fax_line(line)])),
        "email": " / ".join(unique(emails)),
        "contact_persons": " / ".join(unique([clean_label(line) for line in lines if is_contact_line(line)])),
    }
    return {key: value or None for key, value in parsed.items()}


def read_agents(workbook_path: Path) -> list[dict[str, str | None]]:
    worksheet = load_workbook(workbook_path, data_only=True)["船代信息"]
    agents: list[dict[str, str | None]] = []
    current_port: str | None = None

    for row in worksheet.iter_rows(values_only=True):
        cells = [str(value).strip() for value in row if value and str(value).strip()]
        if not cells:
            continue

        if len(cells) == 1 and "\n" not in cells[0] and not re.search(r"\b(agency|shipping|ocean|co\.|ltd)\b", cells[0], re.I):
            current_port = cells[0]
            continue

        if not current_port:
            continue

        for raw_text in cells:
            parsed = parse_agent(raw_text)
            parsed["port_name"] = current_port
            parsed["raw_text"] = clean_raw_text(raw_text)
            agents.append(parsed)

    for index, agent in enumerate(agents, start=1):
        agent["code"] = f"SA{index:03d}"
    return agents


def sql_literal(value: object) -> str:
    if value is None or str(value) == "":
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def build_shipping_agents_sql(agents: list[dict[str, str | None]]) -> str:
    value_lines = []
    for agent in agents:
        values = [
            agent["code"],
            agent["port_name"],
            agent["agency_name"],
            agent["address"],
            agent["tel"],
            agent["fax"],
            agent["email"],
            agent["contact_persons"],
            agent["raw_text"],
        ]
        value_lines.append("  (" + ", ".join(sql_literal(value) for value in values) + ")")
    joined_values = ",\n".join(value_lines)

    return f"""with module as (
  select id from public.database_modules where slug = '{MODULE_SLUG}'
),
rows (code, port_name, agency_name, address, tel, fax, email, contact_persons, raw_text) as (
values
{joined_values}
)
insert into public.shipping_agents (
  module_id, code, port_name, agency_name, address, tel, fax, email,
  contact_persons, raw_text, source_document
)
select
  module.id, rows.code, rows.port_name, rows.agency_name, rows.address, rows.tel,
  rows.fax, rows.email, rows.contact_persons, rows.raw_text, '{SOURCE_DOCUMENT}'
from rows
cross join module
on conflict (code) do update
set
  module_id = excluded.module_id,
  port_name = excluded.port_name,
  agency_name = excluded.agency_name,
  address = excluded.address,
  tel = excluded.tel,
  fax = excluded.fax,
  email = excluded.email,
  contact_persons = excluded.contact_persons,
  raw_text = excluded.raw_text,
  source_document = excluded.source_document,
  updated_at = now();
"""


def replace_shipping_agents_block(seed_path: Path, block: str) -> None:
    text = seed_path.read_text(encoding="utf-8")
    marker = "with module as (\n  select id from public.database_modules where slug = 'port_agency_info'\n),\nrows (code, port_name, agency_name, address, tel, fax, email, contact_persons, raw_text) as ("
    start = text.index(marker)
    seed_path.write_text(text[:start] + block, encoding="utf-8")


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: generate_shipping_agents_seed.py <中国港口信息.xlsx> <seed_sql_path>")

    workbook_path = Path(sys.argv[1])
    seed_path = Path(sys.argv[2])
    agents = read_agents(workbook_path)
    replace_shipping_agents_block(seed_path, build_shipping_agents_sql(agents))
    print(f"wrote {len(agents)} shipping agent rows to {seed_path}")


if __name__ == "__main__":
    main()
